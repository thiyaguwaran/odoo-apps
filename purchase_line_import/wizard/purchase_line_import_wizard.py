# -*- coding: utf-8 -*-
import io
import binascii
import openpyxl
from odoo import models, fields, api, _
from odoo.exceptions import ValidationError

class PurchaseOrderLineImportWizard(models.TransientModel):
    _name = 'purchase.order.line.import.wizard'
    _description = 'Import Purchase Order Lines'

    data_file = fields.Binary(string='Excel File', required=True)
    file_name = fields.Char(string='Filename')

    def action_import(self):
        self.ensure_one()
        active_id = self.env.context.get('active_id')
        if not active_id:
            return False

        order = self.env['purchase.order'].browse(active_id)
        if not order.exists():
            return False

        if not self.data_file:
            raise ValidationError(_("Please upload an Excel file."))

        # Decode the file
        try:
            file_data = binascii.a2b_base64(self.data_file)
            f = io.BytesIO(file_data)
            wb = openpyxl.load_workbook(f, data_only=True)
            sheet = wb.active
        except Exception as e:
            raise ValidationError(_("Invalid file format! Please upload a valid Excel (.xlsx) file. Error: %s") % str(e))

        if sheet.max_row <= 1:
            raise ValidationError(_("The uploaded file is empty or has no data rows."))

        # Parse headers
        headers = [str(cell.value).strip().lower() if cell.value is not None else '' for cell in sheet[1]]
        
        # Helper to find column index based on name variations
        def get_col_index(aliases):
            for alias in aliases:
                normalized = alias.strip().lower()
                if normalized in headers:
                    return headers.index(normalized)
            return -1

        part_num_idx = get_col_index(['part number', 'part_number', 'partno', 'part no', 'internal reference', 'default_code', 'product'])
        qty_idx = get_col_index(['quantity', 'qty', 'product_qty', 'count'])
        price_idx = get_col_index(['unit price', 'price_unit', 'price', 'rate'])
        desc_idx = get_col_index(['description', 'name', 'product description'])
        uom_idx = get_col_index(['uom', 'unit of measure', 'product_uom'])
        analytic_idx = get_col_index(['analytic account', 'analytic_account', 'analytic', 'analytic account code', 'account_analytic_id'])
        tax_idx = get_col_index(['taxes', 'tax', 'product_taxes', 'taxes_id'])

        if part_num_idx == -1:
            raise ValidationError(_("Could not find a 'Part Number' or 'Internal Reference' column in the Excel file."))

        errors = []
        lines_to_create = []

        # Read data rows
        for row_num in range(2, sheet.max_row + 1):
            row = sheet[row_num]
            
            # Check if Part Number is present
            part_number_cell = row[part_num_idx].value
            if part_number_cell is None or str(part_number_cell).strip() == '':
                # Skip completely empty rows
                if all(cell.value is None for cell in row):
                    continue
                errors.append(_("Row %d: Missing Part Number.") % row_num)
                continue

            part_number = str(part_number_cell).strip()

            # Find product
            product = self.env['product.product'].search([
                '|', ('default_code', '=', part_number), ('barcode', '=', part_number)
            ], limit=1)

            if not product:
                errors.append(_("Row %d: Product not found for Part Number '%s'.") % (row_num, part_number))
                continue

            # Parse quantity
            quantity = 1.0
            if qty_idx != -1:
                qty_val = row[qty_idx].value
                if qty_val is not None:
                    try:
                        quantity = float(qty_val)
                    except ValueError:
                        errors.append(_("Row %d: Invalid quantity '%s'. Must be a number.") % (row_num, str(qty_val)))
                        continue

            # Parse price
            price_unit = None
            if price_idx != -1:
                price_val = row[price_idx].value
                if price_val is not None and str(price_val).strip() != '':
                    try:
                        price_unit = float(price_val)
                    except ValueError:
                        errors.append(_("Row %d: Invalid unit price '%s'. Must be a number.") % (row_num, str(price_val)))
                        continue

            # Parse Description
            description = None
            if desc_idx != -1:
                desc_val = row[desc_idx].value
                if desc_val is not None:
                    description = str(desc_val).strip()

            # Parse UOM
            uom = None
            if uom_idx != -1:
                uom_val = row[uom_idx].value
                if uom_val is not None and str(uom_val).strip() != '':
                    uom_str = str(uom_val).strip()
                    uom = self.env['uom.uom'].search([('name', '=ilike', uom_str)], limit=1)
                    if not uom:
                        errors.append(_("Row %d: UoM '%s' not found in system.") % (row_num, uom_str))
                        continue

            # Parse Analytic Account
            analytic_account = None
            if analytic_idx != -1:
                analytic_val = row[analytic_idx].value
                if analytic_val is not None and str(analytic_val).strip() != '':
                    analytic_str = str(analytic_val).strip()
                    analytic_account = self.env['account.analytic.account'].search([
                        '|', ('code', '=ilike', analytic_str), ('name', '=ilike', analytic_str)
                    ], limit=1)
                    if not analytic_account:
                        errors.append(_("Row %d: Analytic Account '%s' not found in system.") % (row_num, analytic_str))
                        continue

            # Parse Taxes
            taxes = None
            if tax_idx != -1:
                tax_val = row[tax_idx].value
                if tax_val is not None and str(tax_val).strip() != '':
                    tax_names = [t.strip() for t in str(tax_val).split(',') if t.strip()]
                    tax_ids = []
                    for tax_name in tax_names:
                        tax_rec = self.env['account.tax'].search([
                            ('name', '=ilike', tax_name),
                            ('type_tax_use', '=', 'purchase'),
                            ('company_id', '=', order.company_id.id)
                        ], limit=1)
                        if not tax_rec:
                            errors.append(_("Row %d: Tax '%s' not found in system.") % (row_num, tax_name))
                        else:
                            tax_ids.append(tax_rec.id)
                    if tax_ids:
                        taxes = [(6, 0, tax_ids)]

            # Prepare values for Odoo 18 create
            try:
                line_vals = {
                    'order_id': order.id,
                    'product_id': product.id,
                }
                if quantity is not None:
                    line_vals['product_qty'] = quantity
                if price_unit is not None:
                    line_vals['price_unit'] = price_unit
                if description:
                    line_vals['name'] = description
                if uom:
                    line_vals['product_uom'] = uom.id
                if analytic_account:
                    line_vals['analytic_distribution'] = {str(analytic_account.id): 100.0}
                if taxes is not None:
                    line_vals['taxes_id'] = taxes
                    
                lines_to_create.append(line_vals)
            except Exception as ex:
                errors.append(_("Row %d: Error preparing line data. Reason: %s") % (row_num, str(ex)))

        # Raise all collected errors
        if errors:
            raise ValidationError("\n".join(errors))

        # Create the purchase order lines
        if lines_to_create:
            self.env['purchase.order.line'].create(lines_to_create)
            order.message_post(body=_("Successfully imported %d lines from Excel file '%s'.") % (len(lines_to_create), self.file_name or ''))

        return {'type': 'ir.actions.act_window_close'}
